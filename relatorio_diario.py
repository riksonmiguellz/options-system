"""
Relatório diário automatizado: escaneia mercado, gera Excel e envia tudo ao Telegram.
Agendado para rodar às 18:50 de segunda a sexta.
"""

import json
import requests
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from scanner import escanear_mercado
from telegram_bot import enviar_mensagem, TELEGRAM_TOKEN, TELEGRAM_CHAT_ID


def enviar_arquivo_telegram(caminho_arquivo: str, caption: str = "") -> bool:
    """Envia arquivo ao Telegram via sendDocument."""
    url = f"https://api.telegram.org/bot{TELEGRAM_TOKEN}/sendDocument"
    try:
        with open(caminho_arquivo, "rb") as f:
            r = requests.post(url, data={
                "chat_id": TELEGRAM_CHAT_ID,
                "caption": caption,
            }, files={"document": (caminho_arquivo.split("/")[-1].split("\\")[-1], f)}, timeout=30)
        return r.status_code == 200
    except Exception as e:
        print(f"Erro ao enviar arquivo: {e}")
        return False


def gerar_excel(df: pd.DataFrame) -> str:
    """Gera Excel formatado e retorna caminho do arquivo."""
    colunas_pt = {
        "ativo": "Ativo", "simbolo": "Símbolo", "nome": "Nome", "tipo": "Tipo",
        "strike": "Strike", "spot": "Spot", "vencimento": "Vencimento",
        "dias_vencimento": "Dias p/ Venc.", "preco_mercado": "Preço Mercado",
        "bid": "Bid", "ask": "Ask", "close": "Fechamento", "volume": "Volume",
        "volume_financeiro": "Vol. Financeiro", "market_maker": "Market Maker",
        "tipo_exercicio": "Tipo Exercício", "lote": "Lote", "variacao": "Variação %",
        "bs_preco": "Black-Scholes", "binomial_preco": "Binomial",
        "monte_carlo_preco": "Monte Carlo", "heston_preco": "Heston",
        "media_modelos": "Média Modelos", "distorcao_bs_pct": "Dist. BS %",
        "distorcao_binomial_pct": "Dist. Binom %", "distorcao_mc_pct": "Dist. MC %",
        "distorcao_heston_pct": "Dist. Heston %", "distorcao_media_pct": "Dist. Média %",
        "iv_ativo_pct": "IV Ativo %", "setor": "Setor",
        "score_oplab": "Score OpLab", "data_analise": "Data Análise",
    }
    df = df.rename(columns=colunas_pt)

    cols_resumo = [
        "Ativo", "Símbolo", "Tipo", "Strike", "Spot", "Dias p/ Venc.",
        "Preço Mercado", "Black-Scholes", "Binomial", "Monte Carlo", "Heston",
        "Média Modelos", "Dist. Média %", "Vol. Financeiro", "IV Ativo %",
        "Setor", "Score OpLab",
    ]

    agora = datetime.now()
    nome_arquivo = f"C:/Users/Miguel/options-system/relatorio_opcoes_{agora.strftime('%Y%m%d_%H%M%S')}.xlsx"

    with pd.ExcelWriter(nome_arquivo, engine="openpyxl") as writer:
        df[cols_resumo].to_excel(writer, sheet_name="Resumo", index=False, startrow=2)
        df.to_excel(writer, sheet_name="Dados Completos", index=False, startrow=2)

        wb = writer.book
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=10)
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        # Aba Resumo
        ws = wb["Resumo"]
        ws.merge_cells("A1:Q1")
        t = ws["A1"]
        t.value = f"RELATÓRIO DE OPÇÕES — {agora.strftime('%d/%m/%Y %H:%M')}"
        t.font = Font(bold=True, size=14, color="FFFFFF")
        t.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        t.alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:Q2")
        s = ws["A2"]
        s.value = "Scanner: >=180 dias | Liquidez >= R$ 1.000/dia | Top 20 por distorção média | 4 Modelos: BS + Binomial + MC + Heston"
        s.font = Font(italic=True, size=10, color="4472C4")
        s.alignment = Alignment(horizontal="center")

        for col in range(1, len(cols_resumo) + 1):
            cell = ws.cell(row=3, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin_border

        for row in range(4, 4 + len(df)):
            for col in range(1, len(cols_resumo) + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")
            dist_cell = ws.cell(row=row, column=13)
            try:
                val = float(dist_cell.value or 0)
                if val >= 95:
                    dist_cell.fill = green_fill
                    dist_cell.font = Font(bold=True, color="006100")
                elif val >= 90:
                    dist_cell.fill = yellow_fill
            except Exception:
                pass

        for col in range(1, len(cols_resumo) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 15
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["P"].width = 30

        # Aba Dados Completos
        ws2 = wb["Dados Completos"]
        ws2.merge_cells("A1:AF1")
        t2 = ws2["A1"]
        t2.value = f"DADOS COMPLETOS — {agora.strftime('%d/%m/%Y %H:%M')}"
        t2.font = Font(bold=True, size=12, color="FFFFFF")
        t2.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        t2.alignment = Alignment(horizontal="center")

        for col in range(1, len(df.columns) + 1):
            cell = ws2.cell(row=3, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            ws2.column_dimensions[get_column_letter(col)].width = 14

    return nome_arquivo


def gerar_resumo_telegram(df: pd.DataFrame) -> str:
    """Gera texto resumo para enviar no Telegram."""
    agora = datetime.now()
    linhas = [
        "🔍 *SCANNER DE OPÇÕES — RELATÓRIO DIÁRIO*",
        f"📅 {agora.strftime('%d/%m/%Y %H:%M')}",
        "🎯 Filtro: ≥180 dias | Liquidez ≥ R$ 1.000/dia",
        f"📊 Top {len(df)} oportunidades",
        "",
    ]

    for i, (_, row) in enumerate(df.iterrows(), 1):
        if i > 10:
            break
        emoji = "🟢" if row["distorcao_media_pct"] >= 20 else "🟡" if row["distorcao_media_pct"] >= 10 else "⚪"
        linhas.append(f"{emoji} *{i}. {row['simbolo']}* ({row['ativo']})")
        linhas.append(f"  {row['tipo']} | Strike {row['strike']:.2f} | Spot {row['spot']:.2f}")
        linhas.append(f"  Venc: {row['vencimento']} ({row['dias_vencimento']}d)")
        linhas.append(f"  Preço mkt: R$ {row['preco_mercado']:.2f} | Média modelos: R$ {row['media_modelos']:.4f}")
        linhas.append(f"  📐 Distorção média: {row['distorcao_media_pct']:.2f}%")
        linhas.append("")

    if len(df) > 10:
        linhas.append(f"📋 *Mais {len(df) - 10} oportunidades no Excel anexo*")

    linhas.append("⚠️ Distorções >95% — verificar liquidez real antes de operar.")
    return "\n".join(linhas)


def executar():
    """Executa o fluxo completo: scan → Excel → Telegram."""
    agora = datetime.now()
    print(f"[{agora}] Iniciando relatório diário...")

    # 1. Escanear mercado
    df = escanear_mercado(dias_min=180, liquidez_min=1000, top_n=20)

    if df.empty:
        msg = f"⚠️ *Relatório {agora.strftime('%d/%m/%Y')}*\nNenhuma oportunidade encontrada com os filtros atuais."
        enviar_mensagem(msg)
        print("Nenhuma oportunidade encontrada.")
        return

    print(f"Encontradas {len(df)} oportunidades.")

    # 2. Gerar Excel
    caminho_excel = gerar_excel(df)
    print(f"Excel salvo: {caminho_excel}")

    # 3. Enviar resumo em texto ao Telegram
    resumo = gerar_resumo_telegram(df)
    ok_msg = enviar_mensagem(resumo)
    print(f"Resumo Telegram: {'OK' if ok_msg else 'FALHA'}")

    # 4. Enviar Excel ao Telegram
    ok_file = enviar_arquivo_telegram(
        caminho_excel,
        caption=f"📊 Relatório de Opções — {agora.strftime('%d/%m/%Y')}\nTop {len(df)} oportunidades por distorção média"
    )
    print(f"Excel Telegram: {'OK' if ok_file else 'FALHA'}")

    print(f"[{datetime.now()}] Relatório concluído.")


if __name__ == "__main__":
    executar()
